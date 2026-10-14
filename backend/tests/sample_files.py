"""Small, valid files for the attachment tests, built in memory: a one-pixel PNG, a PDF with real
text objects (hand-assembled with a correct xref so pypdf reads it without repair), an encrypted
PDF, and a DOCX from python-docx."""

import base64
from collections.abc import Sequence
from io import BytesIO

from docx import Document
from pypdf import PdfReader, PdfWriter

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def minimal_pdf(*pages: str) -> bytes:
    """A PDF whose pages each show one line of ``pages`` in Helvetica."""
    objects: list[bytes] = []
    page_ids = [4 + 2 * index for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for page_id, text in zip(page_ids, pages, strict=True):
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 24 Tf 72 712 Td ({escaped}) Tj ET".encode("latin-1")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {page_id + 1} 0 R >>".encode()
        )
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    out = BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode())
    return out.getvalue()


def encrypted_pdf() -> bytes:
    """A PDF that needs a user password to open (so an empty-password decrypt fails)."""
    writer = PdfWriter(clone_from=PdfReader(BytesIO(minimal_pdf("Top secret"))))
    writer.encrypt("locked")
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def docx_bytes(paragraphs: Sequence[str], table: Sequence[Sequence[str]] = ()) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, cells in enumerate(table):
            for column_index, text in enumerate(cells):
                grid.cell(row_index, column_index).text = text
    out = BytesIO()
    document.save(out)
    return out.getvalue()
