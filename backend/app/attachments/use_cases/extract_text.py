"""Text extraction for uploaded files, run once at upload time (never by the worker)."""

from __future__ import annotations

import asyncio
from collections.abc import Callable
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.attachments.models import AttachmentKind
from app.core.config import attachment_settings

NO_TEXT = "No extractable text"
ENCRYPTED = "Encrypted PDF"


@dataclass(frozen=True, slots=True)
class Extraction:
    """What extraction yielded: the (capped) text, the uncapped character count, or a short reason."""

    text: str | None
    text_chars: int | None
    error: str | None


# Each extractor returns the text, or ``None`` with the reason it could not be read.
type Extracted = tuple[str | None, str | None]


def _pdf_text(data: bytes) -> Extracted:
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        return None, ENCRYPTED
    return "\f".join((page.extract_text() or "") for page in reader.pages), None


def _docx_text(data: bytes) -> Extracted:
    document = Document(BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(lines), None


def _plain_text(data: bytes) -> Extracted:
    return data.decode("utf-8", errors="replace").replace("\x00", ""), None


_EXTRACTORS: dict[AttachmentKind, Callable[[bytes], Extracted]] = {
    AttachmentKind.PDF: _pdf_text,
    AttachmentKind.DOCX: _docx_text,
    AttachmentKind.TEXT: _plain_text,
}


def _extract(kind: AttachmentKind, data: bytes) -> Extraction:
    """Synchronous and CPU-bound: called through ``asyncio.to_thread``."""
    try:
        text, error = _EXTRACTORS[kind](data)
    except Exception as failure:  # noqa: BLE001 — a parser's failure on a corrupt file must not fail the upload
        text, error = None, type(failure).__name__
    if text is None:
        return Extraction(text=None, text_chars=None, error=(error or NO_TEXT)[:255])
    if kind != AttachmentKind.TEXT and not text.strip():
        return Extraction(text=None, text_chars=0, error=NO_TEXT)
    return Extraction(text=text[: attachment_settings.attachment_max_extracted_chars], text_chars=len(text), error=None)


class ExtractText:
    """The text of a PDF (page texts joined with form feeds), a DOCX (paragraphs and table cells)
    or a text file (UTF-8 with replacement, NULs stripped), capped at
    ``ATTACHMENT_MAX_EXTRACTED_CHARS``. Images have no text. A file the parser cannot read gives
    an ``error`` instead of failing — the upload still succeeds and the model is told the
    attachment was unreadable."""

    def __init__(self) -> None:
        pass

    async def execute(self, *, kind: AttachmentKind, data: bytes) -> Extraction:
        if kind == AttachmentKind.IMAGE:
            return Extraction(text=None, text_chars=None, error=None)
        return await asyncio.to_thread(_extract, kind, data)
